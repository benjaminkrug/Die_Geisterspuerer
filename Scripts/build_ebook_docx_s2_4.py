"""
Baut ein KDP-fertiges eBook-DOCX fuer Staffel 2, Band 4 ("Der Park, der nicht
vergisst") aus Manuskript_S2-4_Komplett.md.

Abgeleitet von build_ebook_docx_s2_3.py — bewusst NICHT neu erfunden. Dieselbe
Mischung wie dort: eBook-Mechanik (Heading-1-Navigation, Bookmark/Hyperlink-TOC,
kein Word-Feld) plus IMPORTIERTE Textaufbereitung aus
build_taschenbuch_docx_s2_4.py (single source of truth fuer ENDE-Regex,
Szenentrenner-Logik, Typografie, Frontmatter, ASIN-Metadaten).

Verwendung:
    python Scripts/build_ebook_docx_s2_4.py            # DOCX bauen
    python Scripts/build_ebook_docx_s2_4.py --pruefen  # nur pruefen (ohne python-docx)


★ UNTERSCHIEDE ZU build_ebook_docx_s2_3.py — jeder mit Grund:

1. KEIN NEXT_BAND-TEASER. S2-5 ist bislang nicht geschrieben —
   build_taschenbuch_docx_s2_4.py definiert deshalb gar keine NEXT_BAND_*-
   Konstanten. Dieselbe Regel wie bei S2-1 vor S2-2, S2-2 vor S2-3 und S2-3
   vor S2-4: ein Teaser wird immer erst nach dem naechsten fertigen
   Manuskript geschrieben und maschinell dagegen geprueft.

2. KEIN QR-CODE, SOLANGE ES KEINE ASIN GIBT. BAND_ASIN ist None (S2-4 ist
   nicht veroeffentlicht) — der Rezensions-Block wird ohne Code/Link gebaut,
   dieselbe Regel wie beim Taschenbuch.

3. SERIENUEBERSICHT ZEIGT S2-1 UND S2-2 ALS KAUFBAR (beide live), S2-3 hat
   weiterhin keine ASIN und wird deshalb (wie in S2-3s eigenem Buch) nicht als
   kaufbar aufgefuehrt. S2-4 selbst wird nicht aufgefuehrt (das ist dieses
   Buch) — identisch zu build_taschenbuch_docx_s2_4.py Abschnitt
   "Serienuebersicht".

Gleich geblieben (aus build_ebook_docx_s2_3.py):
  - Kapitel als "Heading 1" statt eigenem zentrierten Absatz (Kindle-Navigation)
  - Inhaltsverzeichnis als direkte Hyperlinks+Bookmarks, kein Word-TOC-Feld
  - Keine Seitengroesse, keine Raender, keine Seitenzahlen, keine Spiegelraender
  - Kein Kapitaelchen-Auftakt: Kindle rendert small-caps unzuverlaessig
  - Widmung und Epigraph stehen wie im Druck unmittelbar VOR Kapitel 1 (hier
    ohnehin leer — Autor-Freigabe steht weiterhin aus)
  - Cross-Promo-Bloecke (Spielbuch, Herrenhaus-Detektive) wie im Taschenbuch
"""

import os
import re
import sys
import importlib.util

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.join(_SCRIPT_DIR, "..")


# ── Textaufbereitung + Frontmatter + Metadaten: EINE Quelle (s. Kopf) ──────────
def _lade_taschenbuch_modul():
    pfad = os.path.join(_SCRIPT_DIR, "build_taschenbuch_docx_s2_4.py")
    spec = importlib.util.spec_from_file_location("_s24_taschenbuch", pfad)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)          # main() ist durch __main__-Guard geschuetzt
    return mod


TB = _lade_taschenbuch_modul()

bereite_body_vor = TB.bereite_body_vor
ereignisfolge = TB.ereignisfolge
apply_typography = TB.apply_typography

AUTHOR = TB.AUTHOR
SERIES_TITLE = TB.SERIES_TITLE
STAFFEL_TITLE = TB.STAFFEL_TITLE
BAND_NUM = TB.BAND_NUM
BAND_TITLE = TB.BAND_TITLE
BAND_SUBTITLE = TB.BAND_SUBTITLE
ERWARTETE_KAPITEL = TB.ERWARTETE_KAPITEL
SCENE_BREAK_SYMBOL = TB.SCENE_BREAK_SYMBOL
BAND_ASIN = TB.BAND_ASIN
Q, E = TB.Q, TB.E
WIDMUNG_ZEILEN = TB.WIDMUNG_ZEILEN
EPIGRAPH_ZEILEN = TB.EPIGRAPH_ZEILEN
EPIGRAPH_SOURCE = TB.EPIGRAPH_SOURCE

INPUT_FILE = TB.INPUT_FILE
QR_IMAGE = TB.QR_IMAGE
OUTPUT_DIR = os.path.join(_ROOT, "Output", "S2-4")
OUTPUT_DOCX = os.path.join(OUTPUT_DIR, "KDP_S2-4_eBook.docx")


# ══════════════════════════════════════════════════════════════════════════════
# DOCX (braucht python-docx)
# ══════════════════════════════════════════════════════════════════════════════

def _docx():
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    return dict(Document=Document, Pt=Pt, Cm=Cm,
                AL=WD_ALIGN_PARAGRAPH, BR=WD_BREAK, qn=qn, Ox=OxmlElement)


D = None


def create_styles(doc):
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = D['Pt'](11)
    pf = style.paragraph_format
    pf.space_before = D['Pt'](0)
    pf.space_after = D['Pt'](0)
    pf.line_spacing = 1.5
    pf.first_line_indent = D['Cm'](0.7)
    pf.alignment = D['AL'].JUSTIFY

    # ★ Kapitel MUESSEN diesen Stil benutzen — daraus baut Kindle die Navigation.
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Georgia"
    h1.font.size = D['Pt'](16)
    h1.font.bold = True
    h1.font.italic = False
    h1.font.color.rgb = None          # erben (schwarz) statt Word-Blau
    h1pf = h1.paragraph_format
    h1pf.alignment = D['AL'].CENTER
    h1pf.first_line_indent = D['Cm'](0)
    h1pf.space_before = D['Pt'](0)
    h1pf.space_after = D['Pt'](0)
    h1pf.keep_with_next = True


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = D['Cm'](0)
    p.add_run().add_break(D['BR'].PAGE)


def add_blank_lines(doc, count=1):
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = D['Cm'](0)
        p.paragraph_format.space_before = D['Pt'](0)
        p.paragraph_format.space_after = D['Pt'](0)


def add_centered_text(doc, text, font_size=11, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = D['AL'].CENTER
    p.paragraph_format.first_line_indent = D['Cm'](0)
    p.paragraph_format.space_before = D['Pt'](0)
    p.paragraph_format.space_after = D['Pt'](0)
    run = p.add_run(text)
    run.font.size = D['Pt'](font_size)
    run.font.name = "Georgia"
    run.bold = bold
    run.italic = italic
    return p


def add_centered_image(doc, image_path, width_inches=1.6):
    """Add a centered image (e.g. QR code). Silently skips if file is missing."""
    if not os.path.isfile(image_path):
        print(f"  WARNUNG: Bild nicht gefunden, wird übersprungen: {image_path}")
        return None
    from docx.shared import Inches
    p = doc.add_paragraph()
    p.alignment = D['AL'].CENTER
    p.paragraph_format.first_line_indent = D['Cm'](0)
    p.paragraph_format.space_before = D['Pt'](0)
    p.paragraph_format.space_after = D['Pt'](0)
    run = p.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    return p


def _add_runs_with_markdown(p, text):
    pattern = r'(\*\*(.+?)\*\*|\*(.+?)\*|([^*]+))'
    for _full, bold_text, italic_text, plain_text in re.findall(pattern, text):
        if bold_text:
            run = p.add_run(bold_text); run.bold = True
        elif italic_text:
            run = p.add_run(italic_text); run.italic = True
        elif plain_text:
            run = p.add_run(plain_text)
        else:
            continue
        run.font.name = "Georgia"
        run.font.size = D['Pt'](11)


def add_formatted_paragraph(doc, text, first_line_indent=True):
    text = apply_typography(text)
    p = doc.add_paragraph()
    if not first_line_indent:
        p.paragraph_format.first_line_indent = D['Cm'](0)
    _add_runs_with_markdown(p, text)
    return p


def _add_bookmark(paragraph, bookmark_id, bookmark_name):
    """Bookmark um den Absatzinhalt (Dokumentstruktur, kein berechnetes Feld —
    jeder DOCX->eBook-Konverter loest das direkt auf)."""
    start = D['Ox']("w:bookmarkStart")
    start.set(D['qn']("w:id"), str(bookmark_id))
    start.set(D['qn']("w:name"), bookmark_name)
    end = D['Ox']("w:bookmarkEnd")
    end.set(D['qn']("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_internal_hyperlink(paragraph, text, anchor, font_size=12):
    """Direkter interner Hyperlink (w:anchor) statt TOC-Feld."""
    hyperlink = D['Ox']("w:hyperlink")
    hyperlink.set(D['qn']("w:anchor"), anchor)

    run = D['Ox']("w:r")
    rPr = D['Ox']("w:rPr")
    rFonts = D['Ox']("w:rFonts")
    rFonts.set(D['qn']("w:ascii"), "Georgia")
    rFonts.set(D['qn']("w:hAnsi"), "Georgia")
    rPr.append(rFonts)
    sz = D['Ox']("w:sz")
    sz.set(D['qn']("w:val"), str(font_size * 2))
    rPr.append(sz)
    run.append(rPr)

    t = D['Ox']("w:t")
    t.set(D['qn']("xml:space"), "preserve")
    t.text = text
    run.append(t)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _kapitel_nummer(title):
    m = re.search(r"Kapitel (\d+)", title)
    if not m:
        raise ValueError(f"Keine Kapitelnummer in Titel gefunden: {title!r}")
    return int(m.group(1))


def add_table_of_contents(doc, chapter_titles):
    """Echtes, klickbares Inhaltsverzeichnis als direkte Hyperlinks zu
    Bookmarks an den Kapitelueberschriften — NICHT als Word-TOC-\\field."""
    add_blank_lines(doc, 3)
    add_centered_text(doc, "Inhalt", font_size=18, bold=True)
    add_blank_lines(doc, 2)

    for title in chapter_titles:
        nummer = _kapitel_nummer(title)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = D['Cm'](0)
        p.paragraph_format.space_after = D['Pt'](6)
        _add_internal_hyperlink(p, title, f"kapitel_{nummer}", font_size=12)

    add_page_break(doc)


def add_front_matter(doc, chapter_titles):
    # Halbtitel
    add_blank_lines(doc, 8)
    add_centered_text(doc, SERIES_TITLE, font_size=18, bold=True)
    add_page_break(doc)

    # Titelseite — mit Reihenname, weil es mehrere Baende mit Nummern gibt
    add_blank_lines(doc, 5)
    add_centered_text(doc, SERIES_TITLE, font_size=22, bold=True)
    add_blank_lines(doc, 1)
    add_centered_text(doc, STAFFEL_TITLE, font_size=13)
    add_blank_lines(doc, 1)
    add_centered_text(doc, BAND_TITLE, font_size=16, italic=True)
    add_blank_lines(doc, 1)
    add_centered_text(doc, f"Band {BAND_NUM}", font_size=12)
    add_blank_lines(doc, 3)
    add_centered_text(doc, AUTHOR, font_size=12)
    add_page_break(doc)

    # Impressum
    add_blank_lines(doc, 12)
    add_centered_text(doc, f"{SERIES_TITLE} – {BAND_TITLE}", font_size=10, bold=True)
    add_centered_text(doc, BAND_SUBTITLE, font_size=10)
    add_blank_lines(doc, 1)
    add_centered_text(doc, f"© 2026 {AUTHOR}", font_size=9)
    add_centered_text(doc, "Alle Rechte vorbehalten.", font_size=9)
    add_blank_lines(doc, 1)
    p = add_centered_text(doc, "", font_size=9)
    run = p.add_run(
        "Dieses Buch ist ein Werk der Fiktion. Namen, Figuren, Orte und "
        "Ereignisse sind frei erfunden. Jede Ähnlichkeit mit tatsächlichen "
        "Personen, lebend oder tot, ist rein zufällig."
    )
    run.font.size = D['Pt'](9)
    run.font.name = "Georgia"
    add_blank_lines(doc, 1)
    add_centered_text(doc, f"Umschlaggestaltung: {AUTHOR}", font_size=9)
    add_centered_text(doc, f"Satz und Layout: {AUTHOR}", font_size=9)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Erstausgabe 2026", font_size=9)
    add_centered_text(doc, "Independently published", font_size=9)
    add_page_break(doc)

    # Inhaltsverzeichnis (Navigationsapparat)
    add_table_of_contents(doc, chapter_titles)

    # Widmung und Epigraph stehen wie im Druck unmittelbar vor Kapitel 1.
    # ⚠️ Beide Listen sind aktuell LEER (Autor-Freigabe steht aus) — die
    # for-Schleifen erzeugen dann einfach keine Zeilen, kein Absturz.
    if WIDMUNG_ZEILEN:
        add_blank_lines(doc, 10)
        for zeile in WIDMUNG_ZEILEN:
            add_centered_text(doc, zeile, font_size=12, italic=True)
            add_blank_lines(doc, 1)
        add_page_break(doc)

    if EPIGRAPH_ZEILEN:
        add_blank_lines(doc, 10)
        for zeile in EPIGRAPH_ZEILEN:
            add_centered_text(doc, zeile, font_size=12, italic=True)
            add_blank_lines(doc, 1)
        add_blank_lines(doc, 1)
        add_centered_text(doc, f"— {EPIGRAPH_SOURCE}", font_size=10, italic=True)


def add_back_matter(doc):
    add_page_break(doc)

    # 1. Rezensions-Bitte + QR (nur mit ASIN — hier keine, s. Dateikopf Punkt 2)
    add_blank_lines(doc, 3)
    add_centered_text(doc, f"{Q}{BAND_TITLE}{E} hat dir gefallen?", font_size=14, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Dann freue ich mich riesig über eine kurze Bewertung auf Amazon "
        "— auch nur ein oder zwei Sätze reichen völlig.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Jede Rezension hilft anderen Kindern (und ihren Eltern), dieses Buch "
        "zu entdecken. Und mir hilft sie, weitere Bände zu schreiben.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 2)

    if BAND_ASIN:
        add_centered_text(doc, "Einfach den Code scannen und eine Bewertung dalassen:",
                          font_size=11, italic=True)
        add_blank_lines(doc, 1)
        add_centered_image(doc, QR_IMAGE, width_inches=1.6)
        add_blank_lines(doc, 1)
        add_centered_text(doc, "(Handykamera auf den Code halten – der Link öffnet sich von selbst.)",
                          font_size=9, italic=True)
        add_blank_lines(doc, 2)

    add_centered_text(doc, "Vielen Dank!", font_size=11)
    add_centered_text(doc, AUTHOR, font_size=11)

    # 2. KEIN TEASER auf S2-5 — nicht geschrieben (s. Dateikopf Punkt 1).

    # 3. Serienuebersicht — nur kaufbare Titel: S2-1 + S2-2 (beide live) + die
    #    fuenf Baende Staffel 1. S2-3 hat weiterhin keine ASIN und wird deshalb
    #    nicht aufgefuehrt. S2-4 selbst wird nicht aufgefuehrt (dieses Buch).
    add_page_break(doc)
    add_blank_lines(doc, 2)
    add_centered_text(doc, f"{SERIES_TITLE} – {STAFFEL_TITLE}", font_size=14, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(doc, "**Band 1:** Der Gast, der blieb", first_line_indent=False)
    add_formatted_paragraph(doc, "**Band 2:** Die Uhr, die rückwärts läuft", first_line_indent=False)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Freier Einstieg — man muss Staffel 1 dafür nicht kennen.",
                      font_size=10, italic=True)

    add_blank_lines(doc, 2)
    add_centered_text(doc, f"{SERIES_TITLE} — die erste Staffel", font_size=14, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Nora, Theo und Schatten hatten schon einmal zu tun. Fünf Fälle, "
        "bevor die Gebundenen anfingen.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    for num, title in [
        (1, "Das Haus, das flüstert"),
        (2, "Der Friedhof ohne Namen"),
        (3, "Schatten sieht mehr"),
        (4, "Die zugemauerte Tür"),
        (5, "Der Schleier"),
    ]:
        add_formatted_paragraph(doc, f"**Band {num}:** {title}", first_line_indent=False)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Man kann sie in jeder Reihenfolge lesen — auch nach diesem Buch.",
                      font_size=10, italic=True)

    # 4. Cross-Verweis (alle beworbenen Titel sind veroeffentlicht)
    add_blank_lines(doc, 1)
    add_centered_text(doc, SCENE_BREAK_SYMBOL, font_size=11)
    add_blank_lines(doc, 1)
    add_centered_text(doc, f"Mehr spannende Abenteuer von {AUTHOR}:", font_size=14, bold=True)
    add_blank_lines(doc, 1)

    add_centered_text(doc, "Die Geisterspürer – als Spielbuch", font_size=13, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Du kennst die Geschichte. Aber was wäre passiert, wenn Nora damals "
        "nicht in den Keller gegangen wäre?",
        first_line_indent=False,
    )
    add_formatted_paragraph(
        doc,
        f"In {Q}Das Haus, das flüstert – Grusel-Spielbuch{E} entscheidest du an "
        "jedem Wendepunkt selbst. Folgst du dem Hund die Treppe hinauf? Gehst "
        "du nachts in den Keller? Liest du das Tagebuch allein — oder mit "
        "Theo?",
        first_line_indent=False,
    )
    add_formatted_paragraph(
        doc,
        "24 verschiedene Enden. Und ein geheimes, das nur die Mutigsten finden.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Grusel-Spielbuch ab 10 Jahren.", font_size=11, italic=True)

    add_blank_lines(doc, 2)
    add_centered_text(doc, "Die Herrenhaus-Detektive", font_size=13, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Niemand darf das alte Herrenhaus betreten. Aber Jonas, Mila und Ben "
        "finden einen Schlüssel — und hinter der verschlossenen Tür wartet "
        "ein Geheimnis, das seit dreißig Jahren niemand lüften durfte.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Spannendes Detektivabenteuer ab 8 Jahren.", font_size=11, italic=True)

    add_blank_lines(doc, 2)
    add_centered_text(doc, "Die Herrenhaus-Detektive – als Spielbuch", font_size=13, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        f"Zwei Fälle, in denen du selbst ermittelst: {Q}Das verbotene "
        f"Herrenhaus{E} und {Q}Das Geheimnis des Brunnens{E}.",
        first_line_indent=False,
    )
    add_formatted_paragraph(
        doc,
        "Welcher Spur folgst du zuerst? Wem glaubst du? Jede Entscheidung "
        "führt dich auf einem anderen Weg durch den Fall — und zu einem "
        "anderen Ende.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Detektiv-Spielbücher ab 8 Jahren.", font_size=11, italic=True)


def add_chapter_heading(doc, title):
    """Echter Heading-1-Stil — daraus baut Kindle die Navigation. Zusaetzlich
    ein Bookmark, auf das die Inhaltsverzeichnis-Hyperlinks zeigen."""
    add_page_break(doc)
    p = doc.add_paragraph(title, style="Heading 1")
    p.paragraph_format.space_before = D['Pt'](36)
    p.paragraph_format.space_after = D['Pt'](24)
    nummer = _kapitel_nummer(title)
    _add_bookmark(p, nummer, f"kapitel_{nummer}")
    return p


def add_scene_break(doc):
    add_blank_lines(doc, 1)
    add_centered_text(doc, SCENE_BREAK_SYMBOL, font_size=11)
    add_blank_lines(doc, 1)


def parse_and_build(doc, content):
    body = bereite_body_vor(content)
    erste_nach_kapitel = False
    erste_nach_trenner = False

    for art, zeile in ereignisfolge(body):
        if art == 'KAPITEL':
            add_chapter_heading(doc, zeile[2:])
            erste_nach_kapitel, erste_nach_trenner = True, False
        elif art == 'TRENNER':
            add_scene_break(doc)
            erste_nach_kapitel, erste_nach_trenner = False, True
        else:
            # Kein Kapitaelchen-Auftakt: Kindle rendert small-caps unzuverlaessig.
            add_formatted_paragraph(
                doc, zeile,
                first_line_indent=not (erste_nach_kapitel or erste_nach_trenner))
            erste_nach_kapitel = erste_nach_trenner = False


# ══════════════════════════════════════════════════════════════════════════════
# PRUEFEN + MAIN
# ══════════════════════════════════════════════════════════════════════════════

def pruefen(content: str) -> int:
    print("=" * 74)
    print("PRUEFUNG DES eBOOK-BILDS — ohne python-docx")
    print("=" * 74)
    fehler = []

    body = bereite_body_vor(content)
    print("ENDE-Marker: entfernt, keine Restzeile ✓")

    ev = ereignisfolge(body)
    kapitel = [z for a, z in ev if a == 'KAPITEL']
    print(f"Kapitel (werden Heading 1 -> Kindle-Navigation): {len(kapitel)} "
          f"(erwartet {ERWARTETE_KAPITEL})")
    if len(kapitel) != ERWARTETE_KAPITEL:
        fehler.append(f"{len(kapitel)} Kapitel statt {ERWARTETE_KAPITEL}")

    print(f"Echte Szenentrenner: {sum(1 for a, _ in ev if a == 'TRENNER')}")
    for i in range(len(ev) - 1):
        if ev[i][0] == 'TRENNER' and ev[i + 1][0] == 'KAPITEL':
            fehler.append(f"Fehl-Ornament vor {ev[i+1][1]!r}")
    print(f"Fehl-Ornamente: {sum(1 for f in fehler if 'Ornament' in f)}")

    print("\nInhaltsverzeichnis-Eintraege, die Kindle bauen wird:")
    for z in kapitel[:3]:
        print(f"   {z[2:]}")
    print(f"   … ({len(kapitel)} insgesamt)")

    aus = apply_typography(body)
    auf, zu = aus.count('„'), aus.count('“')
    gerade = aus.count(chr(34))
    print(f"\nAnfuehrungszeichen nach Typografie: {auf} oeffnend / {zu} schliessend, "
          f"{gerade} gerade uebrig")
    if auf != zu or gerade:
        fehler.append(f"Anfuehrungszeichen unausgeglichen ({auf}/{zu}, {gerade} gerade)")

    print(f"\nFrontmatter (importiert):")
    print(f"   Widmung : {'(leer — Autor-Freigabe steht aus)' if not WIDMUNG_ZEILEN else WIDMUNG_ZEILEN[0]}")
    print(f"   Epigraph: {'(leer — Autor-Freigabe steht aus)' if not EPIGRAPH_ZEILEN else EPIGRAPH_ZEILEN[0]}")

    if BAND_ASIN:
        print(f"\n✅ ASIN vorhanden ({BAND_ASIN}) — Rezensions-Block + QR-Code aktiv.")
        if not os.path.isfile(QR_IMAGE):
            fehler.append(f"QR-Bild fehlt: {QR_IMAGE}")
    else:
        print("\n⚠️  KEINE ASIN — die Rezensions-Seite wird OHNE Link/QR-Code gebaut.")

    print("✅ Kein Teaser (S2-5 ist noch nicht geschrieben).")

    print()
    if fehler:
        print("!! BEFUNDE:")
        for f in dict.fromkeys(fehler):
            print(f"   - {f}")
        return 1
    print("Alle Pruefungen bestanden.")
    return 0


def main():
    global D

    with open(INPUT_FILE, "r", encoding="utf-8") as f:
        content = f.read()

    if "--pruefen" in sys.argv:
        sys.exit(pruefen(content))

    try:
        D = _docx()
    except ImportError:
        print("!! python-docx ist nicht installiert — DOCX kann nicht gebaut werden.")
        print("   Die Textpruefung laeuft trotzdem:")
        print("       python Scripts/build_ebook_docx_s2_4.py --pruefen")
        sys.exit(2)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print(f"Lese: {INPUT_FILE}")

    chapter_titles = re.findall(r"^# (Kapitel \d+.*)$", content, re.MULTILINE)
    if not chapter_titles:
        raise ValueError("Keine Kapitelueberschriften im Manuskript gefunden")

    doc = D['Document']()
    create_styles(doc)
    add_front_matter(doc, chapter_titles)
    parse_and_build(doc, content)
    add_back_matter(doc)

    doc.save(OUTPUT_DOCX)
    print(f"DOCX erstellt: {OUTPUT_DOCX}")

    kapitel = len(re.findall(r'^# Kapitel ', content, re.MULTILINE))
    print(f"Kapitel: {kapitel}")
    if kapitel != ERWARTETE_KAPITEL:
        print(f"!! ABBRUCH-WARNUNG: {kapitel} Kapitel, erwartet {ERWARTETE_KAPITEL}.")
        sys.exit(1)

    print("Frontmatter/Metadaten aus build_taschenbuch_docx_s2_4.py uebernommen —")
    print("eine Quelle fuer Taschenbuch und eBook.")
    if BAND_ASIN:
        print(f"✅ Rezensions-Link + QR-Code aktiv (ASIN {BAND_ASIN}).")
    else:
        print("⚠️  OHNE Rezensions-Link/QR-Code gebaut — keine ASIN hinterlegt.")
    print("✅ Kein Teaser eingebaut (S2-5 ist noch nicht geschrieben).")


if __name__ == "__main__":
    main()
