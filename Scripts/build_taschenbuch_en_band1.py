"""
Baut das KDP-Taschenbuch (Innenteil, DOCX + PDF) der englischen Ausgabe von Band 1:
The Ghost Listeners, Book 1: The Whispering Window.

Format 5,5 x 8,5 Zoll — gilt fuer alle englischen Baende (Entscheidung des Autors 2026-10-17,
Dokumentation/EN_Bibel/A10_Titel_Markt.md).

Inhalt (Kapitel, Rahmentexte, Leseprobe Book 2) kommt aus Scripts/build_ebook_en_band1.py, damit
eBook und Taschenbuch nie auseinanderlaufen. Der Satz folgt dem deutschen Taschenbuch
(Scripts/build_taschenbuch_docx.py): Georgia 11 pt, 1,5 Zeilen, Spiegelraender, Seitenzahlen
ab dem ersten Kapitel, Ornament-Trenner, Kapitel auf neuer Seite. Kein Inhaltsverzeichnis
(wie im deutschen Taschenbuch). Bewertungsbitte ohne QR, bis die ASIN feststeht.

Hochgeladen wird das PDF: Seitenzahl bei KDP = Seitenzahl des PDFs -> daraus der Buchruecken
(Scripts/build_cover.py 1en).

Aufruf:  python Scripts/build_taschenbuch_en_band1.py
Ausgabe: Output/Band1_EN/KDP_Book1_Paperback_EN.docx und .pdf
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_ebook_en_band1 as eb  # noqa: E402  (Inhalt + Rahmen, gemeinsam mit dem eBook)

ROOT = eb.ROOT
OUT_DIR = os.path.join(ROOT, "Output", "Band1_EN")
OUTPUT_DOCX = os.path.join(OUT_DIR, "KDP_Book1_Paperback_EN.docx")
OUTPUT_PDF = os.path.join(OUT_DIR, "KDP_Book1_Paperback_EN.pdf")
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

TRIM_W, TRIM_H = 5.5, 8.5
SCENE_BREAK_SYMBOL = "\u2726  \u2726  \u2726"   # wie im deutschen Taschenbuch


# --- Seite ------------------------------------------------------------------------------------
def setup_page(section):
    """5,5 x 8,5 Zoll, Raender wie das deutsche Taschenbuch (ueber den KDP-Mindestwerten)."""
    section.page_width = Inches(TRIM_W)
    section.page_height = Inches(TRIM_H)
    section.left_margin = Inches(0.75)     # innen (Spiegelraender)
    section.right_margin = Inches(0.5)     # aussen
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.625)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0.3)


def setup_mirror_margins(doc):
    doc.settings.element.append(OxmlElement("w:mirrorMargins"))


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for art, text in (("begin", None), ("instr", " PAGE "), ("end", None)):
        run = p.add_run()
        run.font.name = "Georgia"
        run.font.size = Pt(10)
        if art == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), art)
        run._r.append(el)


def create_styles(doc):
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0.5)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True
    # Englische Silbentrennung im Blocksatz (sonst grosse Wortluecken)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    style.element.get_or_add_rPr().append(lang)
    auto = OxmlElement("w:autoHyphenation")
    doc.settings.element.append(auto)


# --- Bausteine (Druck) -----------------------------------------------------------------------
def add_scene_break(doc):
    eb.add_blank_lines(doc, 1)
    eb.add_centered_text(doc, SCENE_BREAK_SYMBOL, font_size=11)
    eb.add_blank_lines(doc, 1)


def add_chapter_heading(doc, title, erste=False):
    """Wie im deutschen Taschenbuch: neue Seite, 14 pt fett, Luft davor und danach."""
    if not erste:
        eb.add_page_break(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(50)
    p.paragraph_format.space_after = Pt(30)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.size = Pt(14)
    run.font.name = "Georgia"
    run.bold = True


def front_matter(doc, r):
    """Schmutztitel, Titelseite, Impressum — ohne Inhaltsverzeichnis (wie DE-Taschenbuch)."""
    eb.add_blank_lines(doc, 8)
    for z in r["Half title"]:
        eb.zeile(doc, z, 18, bold=True)
    eb.add_page_break(doc)

    eb.add_blank_lines(doc, 5)
    groessen = [14, 12, 24, 12]
    for i, z in enumerate(r["Title page"]):
        eb.zeile(doc, z, groessen[min(i, len(groessen) - 1)], bold=(i == 2))
        eb.add_blank_lines(doc, 2 if i == len(r["Title page"]) - 2 else 1)
    eb.add_page_break(doc)

    eb.add_blank_lines(doc, 6)
    for z in r["Copyright page"]:
        eb.zeile(doc, z, 9)
        eb.add_blank_lines(doc, 1)


def body(doc, content, ende_marker="**END OF BOOK 1**"):
    """Kapitel-Logik wie eb.body: Trenner vor einer Ueberschrift entfaellt (Separator-Bug),
    Ende-Marker raus; dazu Druck-Ueberschrift und Ornament-Trenner."""
    content = content.replace(ende_marker, "")
    content = re.sub(r"(\n\s*---\s*)+$", "\n", content.rstrip() + "\n")
    lines = [l.strip() for l in content.split("\n")]
    erster, kapitel = False, 0
    for i, line in enumerate(lines):
        if not line:
            continue
        if line.startswith("# "):
            add_chapter_heading(doc, line[2:], erste=(kapitel == 0))
            kapitel += 1
            erster = True
            continue
        if line == "---":
            rest = [l for l in lines[i + 1:] if l]
            if not rest or rest[0].startswith("# "):
                continue
            add_scene_break(doc)
            erster = True
            continue
        eb.add_formatted_paragraph(doc, line, first_line_indent=not erster)
        erster = False
    return kapitel


# --- Pruefung am DOCX ------------------------------------------------------------------------
def pruefe_docx(doc, content, kapitel):
    text = "\n".join(p.text for p in doc.paragraphs)
    fehlen = []
    for a in content.split("\n\n"):
        a = a.strip()
        if not a or a in ("---", "**END OF BOOK 1**"):
            continue
        a = a.lstrip("# ").replace("*", "")
        if a not in text:
            fehlen.append(a[:60])
    if fehlen:
        raise ValueError(f"{len(fehlen)} Absaetze fehlen, z. B. {fehlen[:3]}")
    if kapitel != 16:
        raise ValueError(f"{kapitel} Kapitel statt 16")
    if "END OF BOOK 1" in text:
        raise ValueError("Ende-Marker steht noch im Dokument")
    if "Coming soon" not in text or "Did you enjoy" not in text:
        raise ValueError("Anhang unvollstaendig")
    auf, zu = text.count("\u201c"), text.count("\u201d")
    if auf != zu:
        raise ValueError(f"Anfuehrungszeichen unausgeglichen: \u201c {auf} / \u201d {zu}")
    if '"' in text:
        raise ValueError("Gerade Anfuehrungszeichen im Text")
    return auf


# --- PDF ---------------------------------------------------------------------------------------
def zu_pdf():
    if not os.path.exists(SOFFICE):
        raise ValueError(f"LibreOffice fehlt: {SOFFICE}")
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", OUT_DIR, OUTPUT_DOCX],
                   check=True, capture_output=True, timeout=300)
    if not os.path.exists(OUTPUT_PDF):
        raise ValueError("PDF wurde nicht erzeugt")
    masse_exakt()


def masse_exakt():
    """LibreOffice rundet die Seitenbreite auf 140 mm (5,512 statt 5,5 Zoll; die deutschen PDFs
    haben dieselbe Abweichung, z. B. 9,016 statt 9 Zoll). Seite auf exakt 5,5 x 8,5 Zoll setzen:
    weg faellt ein leerer Streifen von 0,3 mm am rechten Rand, der Text liegt >= 12 mm davon."""
    import fitz
    d = fitz.open(OUTPUT_PDF)
    soll = fitz.Rect(0, 0, TRIM_W * 72, TRIM_H * 72)
    for p in d:
        r = p.rect
        if r.width + 0.01 < soll.width or r.height + 0.01 < soll.height:
            raise ValueError(f"PDF-Seite kleiner als das Format: {r.width/72:.3f} x {r.height/72:.3f}")
        p.set_mediabox(soll)
    # Gerade Seitenzahl: KDP haengt sonst selbst eine Leerseite an -- so ist die Zahl im PDF
    # dieselbe wie bei KDP, und der Buchruecken laesst sich exakt berechnen.
    if d.page_count % 2:
        d.new_page(width=soll.width, height=soll.height)
    tmp = OUTPUT_PDF + ".tmp"
    d.save(tmp, garbage=3, deflate=True)
    d.close()
    os.replace(tmp, OUTPUT_PDF)


def pruefe_pdf():
    """Am fertigen Artefakt: Seitengroesse, Seitenzahl, eingebettete Schriften, Kapitel im Text."""
    import fitz
    d = fitz.open(OUTPUT_PDF)
    seiten = d.page_count
    falsch = [i + 1 for i, p in enumerate(d) if abs(p.rect.width / 72 - TRIM_W) > 0.01
              or abs(p.rect.height / 72 - TRIM_H) > 0.01]
    if falsch:
        raise ValueError(f"Seitengroesse falsch auf Seiten {falsch[:5]}")
    nicht_eingebettet = set()
    for p in d:
        for f in p.get_fonts(full=True):
            # (xref, ext, type, basefont, ...) -- ext "n/a" = Schrift nicht eingebettet
            if f[1] == "n/a":
                nicht_eingebettet.add(f[3])
    gesamt = "".join(p.get_text() for p in d)
    # Titel und Reihenname duerfen nirgends getrennt sein
    for wort in ("Whispering", "Window", "Listeners"):
        if re.search(wort[:4] + r"\w*-\s*\n", gesamt) and wort not in gesamt.replace("-\n", ""):
            raise ValueError(f"{wort} ist getrennt")
    seite1bis3 = "".join(d[i].get_text() for i in range(3))
    if "-\n" in seite1bis3:
        raise ValueError("Silbentrennung in der Titelei")
    for n in range(1, 17):
        if f"Chapter {n}:" not in gesamt:
            raise ValueError(f"Chapter {n} fehlt im PDF")
    d.close()
    if nicht_eingebettet:
        raise ValueError(f"Schriften nicht eingebettet: {sorted(nicht_eingebettet)}")
    return seiten


def main():
    r = eb.rahmen()
    content = eb.kapitel_text()

    doc = Document()
    s_front = doc.sections[0]
    setup_page(s_front)
    s_front.header.is_linked_to_previous = False
    s_front.footer.is_linked_to_previous = False
    setup_mirror_margins(doc)
    create_styles(doc)
    cp = doc.core_properties
    cp.title = "The Whispering Window (The Ghost Listeners, Book 1)"
    cp.author = cp.last_modified_by = "Benjamin Krug"
    cp.language = "en-US"

    front_matter(doc, r)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    s_body = doc.sections[-1]
    setup_page(s_body)
    add_page_number_footer(s_body)
    s_body.header.is_linked_to_previous = False

    kapitel = body(doc, content)
    eb.add_scene_break = add_scene_break          # Druck-Ornament auch in der Leseprobe
    eb.back_matter(doc, r)

    # Silbentrennung nur im Fliesstext: zentrierte Zeilen (Titel, Ueberschriften, Impressum,
    # Anhang) nie trennen -- sonst stand auf der Titelseite "The Whispering Win-dow".
    for p in doc.paragraphs:
        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            p._p.get_or_add_pPr().append(OxmlElement("w:suppressAutoHyphens"))

    zeichen = pruefe_docx(doc, content, kapitel)
    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(f"{OUTPUT_DOCX} erstellt — {kapitel} Kapitel, Anfuehrungszeichen {zeichen}/{zeichen}")

    zu_pdf()
    seiten = pruefe_pdf()
    print(f"{OUTPUT_PDF} erstellt — {seiten} Seiten, {TRIM_W} x {TRIM_H} Zoll, Schriften eingebettet")
    print(f"  -> Cover: BAENDE['1en'] in Scripts/build_cover.py auf seiten={seiten} pruefen, "
          f"dann python Scripts/build_cover.py 1en")


if __name__ == "__main__":
    try:
        main()
    except ValueError as e:
        print("FEHLER:", e)
        sys.exit(1)
