"""
Baut das KDP-eBook (DOCX) der englischen Ausgabe von Band 1: The Ghost Listeners, Book 1:
The Whispering Window.

Quellen (alle in Band1/Manuskript_EN/):
  Kapitel_01.md … Kapitel_16.md      — geprüfter Text (pruefe_en.py: 0 Befunde)
  Rahmen/Rahmen_EN.md                — Titelei, Impressum, Anhang-Texte (je Seite ein „## "-Abschnitt)
  Rahmen/EN_Leseprobe_Band2.md       — Sneak peek Book 2

Die Hilfsfunktionen (Stile, Inhaltsverzeichnis-Feld, Kapitelüberschrift als Heading 1, Links)
sind wörtlich aus Scripts/build_ebook_docx.py (Branch Band 1) übernommen, damit beide Ausgaben
gleich aussehen. Der Bau bricht ab, wenn ein Abschnitt fehlt, ein Platzhalter stehen geblieben ist
oder ein Kapiteltext nicht vollständig im Dokument landet.

Aufruf:  python3 Scripts/build_ebook_en_band1.py
Ausgabe: Output/Band1_EN/KDP_Book1_eBook_EN.docx
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EN = os.path.join(ROOT, "Band1", "Manuskript_EN")
RAHMEN = os.path.join(EN, "Rahmen", "Rahmen_EN.md")
LESEPROBE = os.path.join(EN, "Rahmen", "EN_Leseprobe_Band2.md")
OUTPUT_FILE = os.path.join(ROOT, "Output", "Band1_EN", "KDP_Book1_eBook_EN.docx")

# Erst nach Veröffentlichung bekannt (Autor-Entscheidung 2026-10-14: Link per Update).
# Dann eintragen, neu bauen, Datei bei KDP ersetzen.
BOOK1_ASIN_EN = None
REVIEW_URL = f"https://www.amazon.com/review/create-review?asin={BOOK1_ASIN_EN}" if BOOK1_ASIN_EN else None

ABSCHNITTE = ["Half title", "Title page", "Copyright page", "Contents", "What happens next?",
              "Sneak peek heading", "After the sneak peek", "Review request", "Series list"]
TOC_TITEL = "Contents"

def create_styles(doc):
    """Set up document styles for eBook formatting."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Georgia"
    font.size = Pt(11)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0.7)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Configure "Heading 1" so chapter titles are recognized by Kindle/KDP
    # for the auto-generated navigation TOC, while still looking like our
    # centered chapter headings. The key part is that chapters USE this style.
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Georgia"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.italic = False
    h1.font.color.rgb = None  # inherit (black) instead of the default blue
    h1pf = h1.paragraph_format
    h1pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1pf.first_line_indent = Cm(0)
    h1pf.space_before = Pt(0)
    h1pf.space_after = Pt(0)
    h1pf.keep_with_next = True


def add_page_break(doc):
    """Add a page break."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def add_blank_lines(doc, count=1):
    """Add empty paragraphs."""
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


def add_centered_text(doc, text, font_size=11, bold=False, italic=False):
    """Add a centered paragraph."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = "Georgia"
    run.bold = bold
    run.italic = italic
    return p


def add_formatted_paragraph(doc, text, first_line_indent=True, alignment=None):
    """Add a paragraph with inline formatting (*italic*, **bold**)."""
    p = doc.add_paragraph()
    if not first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if alignment:
        p.alignment = alignment
        p.paragraph_format.first_line_indent = Cm(0)

    # Parse inline markdown formatting
    # Handle **bold** and *italic*
    pattern = r'(\*\*(.+?)\*\*|\*(.+?)\*|([^*]+))'
    parts = re.findall(pattern, text)

    for full_match, bold_text, italic_text, plain_text in parts:
        if bold_text:
            run = p.add_run(bold_text)
            run.bold = True
            run.font.name = "Georgia"
            run.font.size = Pt(11)
        elif italic_text:
            run = p.add_run(italic_text)
            run.italic = True
            run.font.name = "Georgia"
            run.font.size = Pt(11)
        elif plain_text:
            run = p.add_run(plain_text)
            run.font.name = "Georgia"
            run.font.size = Pt(11)

    return p


def add_table_of_contents(doc):
    """Insert a real Word TOC field on its own page.

    KDP/Kindle reads this field (built from the Heading 1 chapter titles) and
    turns it into a clickable in-book table of contents. The field needs to be
    updated once in Word/LibreOffice (or by KDP on conversion) to fill in the
    page numbers, but the entries and links are generated automatically.
    """
    add_blank_lines(doc, 3)
    add_centered_text(doc, TOC_TITEL, font_size=18, bold=True)
    add_blank_lines(doc, 2)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()

    # Build the TOC field: \o "1-1" = use Heading 1 only,
    # \h = hyperlinks, \z = hide tab leader/page no in web view, \u = outline level
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-1" \\h \\z \\u'

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    # Placeholder text shown until the field is updated
    placeholder = OxmlElement("w:t")
    placeholder.text = "Table of contents (updates when opened)."

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r_elem = run._r
    r_elem.append(fld_begin)
    r_elem.append(instr)
    r_elem.append(fld_sep)
    r_elem.append(placeholder)
    r_elem.append(fld_end)

    # Tell Word to update fields automatically when the document is opened,
    # so the TOC and its links are populated without manual interaction.
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    add_page_break(doc)


def add_hyperlink(paragraph, url, text, font_size=11):
    """Klickbarer Link (Kindle oeffnet ihn im Browser). python-docx kann das nicht direkt."""
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None), ("w:u", "single"), ("w:sz", str(font_size * 2))):
        el = OxmlElement(tag)
        if tag == "w:rFonts":
            el.set(qn("w:ascii"), "Georgia")
            el.set(qn("w:hAnsi"), "Georgia")
        else:
            el.set(qn("w:val"), val)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return link


def add_centered_link(doc, url, text, font_size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    add_hyperlink(p, url, text, font_size)
    return p


def add_chapter_heading(doc, title):
    """Add a chapter heading with page break.

    Uses the real "Heading 1" style so Kindle/KDP picks the chapter up for the
    auto-generated navigation TOC and the in-book TOC field resolves to it.
    """
    add_page_break(doc)
    add_blank_lines(doc, 3)
    p = doc.add_paragraph(title, style="Heading 1")
    # Ensure the run font matches our book look (style sets this too, but be safe)
    for run in p.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(16)
        run.bold = True
    add_blank_lines(doc, 2)


def add_scene_break(doc):
    """Add a scene break (centered asterisks)."""
    add_blank_lines(doc, 1)
    add_centered_text(doc, "* * *", font_size=11)
    add_blank_lines(doc, 1)




# --- Festes, verlinktes Inhaltsverzeichnis --------------------------------------------------------
# Das Word-Feld (add_table_of_contents) enthält bis zur Aktualisierung nur einen Platzhalter.
# Ob der KDP-Konverter Felder aktualisiert, ist nicht gesichert -> feste Einträge mit internen
# Links auf Textmarken an den Kapitelüberschriften. Funktioniert ohne Aktualisierung (Nachprüfung
# 2026-10-14).
_bm_id = [0]


def add_bookmark(paragraph, name):
    _bm_id[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bm_id[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bm_id[0]))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, anchor, text, font_size=11):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Georgia")
    fonts.set(qn("w:hAnsi"), "Georgia")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(font_size * 2))
    rpr.append(fonts)
    rpr.append(size)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def kapitel_titel(content):
    return [l[2:].strip() for l in content.split("\n") if l.startswith("# Chapter ")]


def add_static_toc(doc, titel_liste, ueberschrift):
    add_blank_lines(doc, 3)
    add_centered_text(doc, ueberschrift, font_size=18, bold=True)
    add_blank_lines(doc, 2)
    for i, titel in enumerate(titel_liste, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_internal_link(p, f"kapitel_{i:02d}", titel)


def add_chapter_heading_bm(doc, title, nummer):
    """Wie add_chapter_heading, zusätzlich Textmarke für das Inhaltsverzeichnis."""
    add_page_break(doc)
    add_blank_lines(doc, 3)
    p = doc.add_paragraph(title, style="Heading 1")
    for run in p.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(16)
        run.bold = True
    add_bookmark(p, f"kapitel_{nummer:02d}")
    add_blank_lines(doc, 2)


# --- Quellen lesen ---------------------------------------------------------------------------
def rahmen():
    text = open(RAHMEN, encoding="utf-8").read()
    teile = re.split(r"^## (.+?)\s*$", text, flags=re.MULTILINE)
    abschnitte = {teile[i].strip(): teile[i + 1] for i in range(1, len(teile), 2)}
    fehlt = [a for a in ABSCHNITTE if a not in abschnitte]
    if fehlt:
        raise ValueError(f"Rahmen_EN.md: Abschnitt fehlt: {fehlt}")
    aus = {}
    for name in ABSCHNITTE:
        zeilen = [z.strip() for z in abschnitte[name].split("\n")]
        zeilen = [z for z in zeilen if z and not z.startswith("<!--") and z != "---"]
        for z in zeilen:
            if re.search(r"TODO|XXX|\[.*?\]|PLACEHOLDER", z):
                raise ValueError(f"Platzhalter in Rahmen_EN.md ({name}): {z}")
        aus[name] = zeilen
    return aus


def kapitel_text():
    teile = []
    for n in range(1, 17):
        t = open(os.path.join(EN, f"Kapitel_{n:02d}.md"), encoding="utf-8").read().strip()
        if not t.startswith(f"# Chapter {n}:"):
            raise ValueError(f"Kapitel_{n:02d}.md beginnt nicht mit '# Chapter {n}:'")
        teile.append(t)
    return "\n\n".join(teile)


# --- Bausteine ----------------------------------------------------------------------------------
def zeile(doc, text, size=11, bold=False, italic=False):
    """Zentrierte Zeile; **fett** / *kursiv* im Markdown werden übernommen."""
    m = re.fullmatch(r"\*\*(.+)\*\*", text)
    if m:
        text, bold = m.group(1), True
    m = re.fullmatch(r"\*(.+)\*", text)
    if m:
        text, italic = m.group(1), True
    if "*" in text:   # gemischt (z. B. Originaltitel kursiv im Satz)
        p = add_formatted_paragraph(doc, text, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs:
            r.font.size = Pt(size)
        return p
    return add_centered_text(doc, text, font_size=size, bold=bold, italic=italic)


def front_matter(doc, r, titel_liste):
    add_blank_lines(doc, 8)
    for z in r["Half title"]:
        zeile(doc, z, 18, bold=True)
    add_page_break(doc)

    add_blank_lines(doc, 5)
    # Reihe · Book One · Buchtitel (größte Zeile, US-Konvention) · Autor
    groessen = [14, 12, 24, 12]
    for i, z in enumerate(r["Title page"]):
        zeile(doc, z, groessen[min(i, len(groessen) - 1)], bold=(i == 2))
        add_blank_lines(doc, 2 if i == len(r["Title page"]) - 2 else 1)
    add_page_break(doc)

    add_blank_lines(doc, 10)
    for z in r["Copyright page"]:
        zeile(doc, z, 9)
        add_blank_lines(doc, 1)
    add_page_break(doc)

    add_static_toc(doc, titel_liste, r["Contents"][0])


def body(doc, content, ende_marker="**END OF BOOK 1**"):
    """Wie parse_and_build im deutschen Skript: Szenentrenner vor einer Überschrift entfällt,
    der Ende-Marker wird entfernt (danach folgt die Teaser-Seite)."""
    content = content.replace(ende_marker, "")
    content = re.sub(r"(\n\s*---\s*)+$", "\n", content.rstrip() + "\n")
    lines = [l.strip() for l in content.split("\n")]
    erster = False
    nummer = 0
    for i, line in enumerate(lines):
        if not line:
            continue
        if line.startswith("# "):
            nummer += 1
            add_chapter_heading_bm(doc, line[2:], nummer)
            erster = True
            continue
        if line == "---":
            rest = [l for l in lines[i + 1:] if l]
            if not rest or rest[0].startswith("# "):
                continue
            add_scene_break(doc)
            erster = True
            continue
        add_formatted_paragraph(doc, line, first_line_indent=not erster)
        erster = False


def seite(doc, zeilen, titel_groesse=16, absatz_ab=80):
    """Erste Zeile als Überschrift, kurze Zeilen zentriert, lange als Absatz."""
    add_page_break(doc)
    add_blank_lines(doc, 3)
    for i, z in enumerate(zeilen):
        if i == 0:
            zeile(doc, z, titel_groesse, bold=True)
            add_blank_lines(doc, 1)
        elif len(z) > absatz_ab:
            add_formatted_paragraph(doc, z, first_line_indent=False)
            add_blank_lines(doc, 1)
        else:
            if z.startswith("The Ghost Listeners ·"):   # Buch-Hinweis vom Teaser absetzen
                add_blank_lines(doc, 2)
            zeile(doc, z, 12)


def back_matter(doc, r):
    seite(doc, r["What happens next?"])

    # Sneak peek: kleine Zeile, dann das Kapitel ohne Seitenumbruch davor
    add_page_break(doc)
    for z in r["Sneak peek heading"]:
        zeile(doc, z, 10, italic=True)
    probe = open(LESEPROBE, encoding="utf-8").read().strip()
    m = re.match(r"# (.+)\n", probe)
    if not m:
        raise ValueError("EN_Leseprobe_Band2.md: Überschrift fehlt")
    p = add_centered_text(doc, m.group(1), font_size=16, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(24)
    erster = True
    for a in [a.strip() for a in probe[m.end():].split("\n\n") if a.strip()]:
        if a == "---":
            add_scene_break(doc)
            erster = True
            continue
        add_formatted_paragraph(doc, a, first_line_indent=not erster)
        erster = False

    seite(doc, r["After the sneak peek"], titel_groesse=12)
    seite(doc, r["Review request"], titel_groesse=14)
    if REVIEW_URL:
        add_centered_link(doc, REVIEW_URL, "Leave a review", font_size=12)
    seite(doc, r["Series list"], titel_groesse=14)


# --- Kontrolle -------------------------------------------------------------------------------------
def pruefe(doc, content):
    """Jeder Kapitelabsatz muss wörtlich (ohne Markdown-Sternchen) im Dokument stehen."""
    text = "\n".join(p.text for p in doc.paragraphs)
    fehlen = []
    for a in content.split("\n\n"):
        a = a.strip()
        if not a or a in ("---", "**END OF BOOK 1**"):
            continue
        a = a.lstrip("# ").replace("*", "")
        if a not in text:
            fehlen.append(a[:60])
    if fehlen:
        raise ValueError(f"{len(fehlen)} Absätze fehlen im Dokument, z. B. {fehlen[:3]}")
    kapitel = sum(1 for p in doc.paragraphs if p.style.name == "Heading 1")
    if kapitel != 16:
        raise ValueError(f"{kapitel} Kapitelüberschriften statt 16")
    xml = doc.element.xml
    for n in range(1, 17):
        if f'w:anchor="kapitel_{n:02d}"' not in xml or f'w:name="kapitel_{n:02d}"' not in xml:
            raise ValueError(f"Inhaltsverzeichnis: Link oder Textmarke für Kapitel {n} fehlt")
    if "updates when opened" in text or "TOC \\o" in xml:
        raise ValueError("Altes Inhaltsverzeichnis-Feld im Dokument")
    if "END OF BOOK 1" in text:
        raise ValueError("Ende-Marker steht noch im Dokument")
    return kapitel


def main():
    r = rahmen()
    content = kapitel_text()
    doc = Document()
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
    create_styles(doc)
    cp = doc.core_properties
    cp.title = "The Whispering Window (The Ghost Listeners, Book 1)"
    cp.author = cp.last_modified_by = "Benjamin Krug"
    cp.language = "en-US"
    titel_liste = kapitel_titel(content)
    front_matter(doc, r, titel_liste)
    body(doc, content)
    back_matter(doc, r)
    kapitel = pruefe(doc, content)
    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    doc.save(OUTPUT_FILE)
    print(f"{OUTPUT_FILE} erstellt — {kapitel} Kapitel, {len(doc.paragraphs)} Absätze")
    if not REVIEW_URL:
        print("  Hinweis: Rezensionsbitte ohne Link (BOOK1_ASIN_EN nach Veröffentlichung eintragen).")


if __name__ == "__main__":
    try:
        main()
    except ValueError as e:
        print("FEHLER:", e)
        sys.exit(1)
