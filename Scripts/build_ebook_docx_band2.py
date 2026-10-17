"""
Build a properly formatted DOCX for KDP eBook upload from Manuskript_Band2_Komplett.md.
Abgeleitet von build_ebook_docx.py (Band 1, 2026-09-16 korrigiert): echte
Hyperlink+Bookmark-TOC statt Word-Feld, QR-Code-Rezension, aktuelle Backmatter
(Spielbuch statt Chrono-Agenten).

Band-2-spezifisch: Das Manuskript ist reines ASCII (keine deutschen Anfuehrungs-
zeichen), anders als Band 1. Deshalb wird hier zusaetzlich apply_typography()
(aus build_taschenbuch_docx_band3.py uebernommen, weil Band 2s EIGENES
Taschenbuch-Skript diese Funktion gar nicht hat) auf jeden Fliesstext-Absatz
UND auf den Teaser angewendet.
"""

import os
import re
from docx import Document
from docx.shared import Pt, Cm, Emu, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AUTHOR = "Benjamin Krug"
INPUT_FILE = os.path.join(os.path.dirname(__file__), "..", "Band2", "Manuskript", "Manuskript_Band2_Komplett.md")
OUTPUT_FILE = os.path.join(os.path.dirname(__file__), "..", "Output", "Band2", "KDP_Band2_eBook.docx")
QR_IMAGE = os.path.join(os.path.dirname(__file__), "..", "Band2", "Cover", "qr_rezension_band2.png")


# ── Typografie: ASCII -> deutsche Anfuehrungszeichen/Gedankenstriche ─────────
# Uebernommen aus build_taschenbuch_docx_band3.py (Band 2s eigenes Taschenbuch-
# Skript hat diese Funktion NICHT -- das war der beim Planen gefundene Bug).

def typo_quotes(text: str) -> str:
    """Wandelt gerade ASCII-Anfuehrungszeichen in deutsche „…" um.

    Heuristik (paarweise, robust): Wir tracken, ob wir uns INNERHALB einer
    Anfuehrung befinden. Das erste " oeffnet, das naechste schliesst, usw.
    Apostroph ' -> '."""
    out = []
    inside = False
    for idx, ch in enumerate(text):
        if ch == '"':
            if not inside:
                out.append('„')
                inside = True
            else:
                out.append('“')
                inside = False
        elif ch == "'":
            prev = text[idx - 1] if idx > 0 else ' '
            nxt = text[idx + 1] if idx + 1 < len(text) else ' '
            if prev.isalpha() and nxt.isalpha():
                out.append('’')
            elif prev.isalpha():
                out.append('’')
            else:
                out.append(ch)
        else:
            out.append(ch)
    return "".join(out)


def typo_dashes(text: str) -> str:
    """Vereinheitlicht Gedankenstriche auf den deutschen Halbgeviertstrich (–)."""
    text = re.sub(r'\s—\s', ' – ', text)
    text = re.sub(r'\s—', ' –', text)
    text = re.sub(r'—\s', '– ', text)
    return text


def apply_typography(text: str) -> str:
    """Alle typografischen Verbesserungen auf einen Textabsatz anwenden."""
    text = typo_quotes(text)
    text = typo_dashes(text)
    return text


def create_styles(doc):
    """Set up document styles for eBook formatting."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Georgia"
    font.size = Pt(11)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0.7)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Georgia"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.italic = False
    h1.font.color.rgb = None
    h1pf = h1.paragraph_format
    h1pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1pf.first_line_indent = Cm(0)
    h1pf.space_before = Pt(0)
    h1pf.space_after = Pt(0)
    h1pf.keep_with_next = True


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def add_blank_lines(doc, count=1):
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


def add_centered_text(doc, text, font_size=11, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = "Georgia"
    run.bold = bold
    run.italic = italic
    return p


def add_centered_image(doc, image_path, width_inches=1.6):
    """Add a centered image (e.g. QR code). Silently skips if file is missing."""
    if not os.path.isfile(image_path):
        print(f"  WARNUNG: Bild nicht gefunden, wird übersprungen: {image_path}")
        return None
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    return p


def add_formatted_paragraph(doc, text, first_line_indent=True, alignment=None):
    """Add a paragraph with inline formatting (*italic*, **bold**)."""
    p = doc.add_paragraph()
    if not first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if alignment:
        p.alignment = alignment
        p.paragraph_format.first_line_indent = Cm(0)

    pattern = r'(\*\*(.+?)\*\*|\*(.+?)\*|([^*]+))'
    parts = re.findall(pattern, text)

    for full_match, bold_text, italic_text, plain_text in parts:
        if bold_text:
            run = p.add_run(bold_text)
            run.bold = True
            run.font.name = "Georgia"
            run.font.size = Pt(11)
        elif italic_text:
            run = p.add_run(italic_text)
            run.italic = True
            run.font.name = "Georgia"
            run.font.size = Pt(11)
        elif plain_text:
            run = p.add_run(plain_text)
            run.font.name = "Georgia"
            run.font.size = Pt(11)

    return p


def _add_bookmark(paragraph, bookmark_id, bookmark_name):
    """Wrap a paragraph's content in a Word bookmark (plain document structure,
    not a computed field -- any DOCX->ebook converter resolves it directly)."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_internal_hyperlink(paragraph, text, anchor, font_size=12):
    """Add a run inside `paragraph` that links to an in-document bookmark
    (w:anchor). Plain structural hyperlink, not a TOC field -- works even
    if the converter does not execute Word field codes."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Georgia")
    rFonts.set(qn("w:hAnsi"), "Georgia")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(font_size * 2))
    rPr.append(sz)
    run.append(rPr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _kapitel_nummer(title):
    """Extract the chapter number from a 'Kapitel N – Titel' string."""
    m = re.search(r"Kapitel (\d+)", title)
    if not m:
        raise ValueError(f"Keine Kapitelnummer in Titel gefunden: {title!r}")
    return int(m.group(1))


def add_table_of_contents(doc, chapter_titles):
    """Insert a real, clickable in-book table of contents (Hyperlink+Bookmark,
    nicht Word-Feld -- siehe build_ebook_docx.py fuer die volle Begruendung)."""
    add_blank_lines(doc, 3)
    add_centered_text(doc, "Inhalt", font_size=18, bold=True)
    add_blank_lines(doc, 2)

    for title in chapter_titles:
        nummer = _kapitel_nummer(title)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(6)
        _add_internal_hyperlink(p, title, f"kapitel_{nummer}", font_size=12)

    add_page_break(doc)


def add_front_matter(doc, chapter_titles):
    """Add title page, half-title, and copyright page."""
    add_blank_lines(doc, 8)
    add_centered_text(doc, "Die Geisterspürer", font_size=18, bold=True)
    add_page_break(doc)

    add_blank_lines(doc, 5)
    add_centered_text(doc, "Die Geisterspürer", font_size=22, bold=True)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Der Friedhof ohne Namen", font_size=16, italic=True)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Band 2", font_size=12)
    add_blank_lines(doc, 3)
    add_centered_text(doc, AUTHOR, font_size=12)
    add_page_break(doc)

    add_blank_lines(doc, 12)
    add_centered_text(doc, "Die Geisterspürer – Der Friedhof ohne Namen", font_size=10, bold=True)
    add_centered_text(doc, "Band 2", font_size=10)
    add_blank_lines(doc, 1)
    add_centered_text(doc, f"© 2026 {AUTHOR}", font_size=9)
    add_centered_text(doc, "Alle Rechte vorbehalten.", font_size=9)
    add_blank_lines(doc, 1)
    p = add_centered_text(doc, "", font_size=9)
    run = p.add_run(
        "Dieses Buch ist ein Werk der Fiktion. Namen, Figuren, Orte und "
        "Ereignisse sind frei erfunden. Jede Ähnlichkeit mit tatsächlichen "
        "Personen, lebend oder tot, ist rein zufällig."
    )
    run.font.size = Pt(9)
    run.font.name = "Georgia"
    add_blank_lines(doc, 1)
    add_centered_text(doc, f"Umschlaggestaltung: {AUTHOR}", font_size=9)
    add_centered_text(doc, f"Satz und Layout: {AUTHOR}", font_size=9)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Erstausgabe 2026", font_size=9)
    add_centered_text(doc, "Independently published", font_size=9)
    add_page_break(doc)

    add_table_of_contents(doc, chapter_titles)


def add_back_matter(doc):
    """Add back matter: review request, Band 3 teaser, series list, cross-references."""
    add_page_break(doc)

    # --- 1. Rezensions-Bitte ---
    add_blank_lines(doc, 3)
    add_centered_text(doc, 'Hat dir „Der Friedhof ohne Namen“ gefallen?', font_size=14, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Dann freue ich mich riesig über eine kurze Bewertung auf Amazon "
        "— auch nur ein oder zwei Sätze reichen völlig.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Jede Rezension hilft anderen Kindern (und ihren Eltern), dieses Buch "
        "zu entdecken. Und mir hilft sie, weitere Bände zu schreiben.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 2)

    add_centered_text(doc, "Einfach den Code scannen und eine Bewertung dalassen:", font_size=11, italic=True)
    add_blank_lines(doc, 1)
    add_centered_image(doc, QR_IMAGE, width_inches=1.6)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "(Handykamera auf den Code halten – der Link öffnet sich von selbst.)", font_size=9, italic=True)

    add_blank_lines(doc, 2)
    add_centered_text(doc, "Vielen Dank!", font_size=11)
    add_centered_text(doc, AUTHOR, font_size=11)

    # --- 2. Teaser Band 3 ---
    # Echter Anfang von Band 3, Kapitel 1 (verifiziert gegen Band3/Manuskript/
    # Kapitel_01.md, 2026-09-16). Wortgleich mit dem bereits korrigierten
    # Taschenbuch-Teaser (build_taschenbuch_docx_band2.py); hier zusaetzlich
    # durch apply_typography() gejagt, weil dort noch ASCII-Anfuehrungszeichen
    # hart codiert sind.
    add_page_break(doc)
    add_blank_lines(doc, 2)
    add_centered_text(doc, "Weiterlesen? Hier kommt eine Vorschau auf Band 3:", font_size=11, italic=True)
    add_blank_lines(doc, 2)
    add_centered_text(doc, "Die Geisterspürer", font_size=18, bold=True)
    add_centered_text(doc, "Schatten sieht mehr", font_size=14, italic=True)
    add_centered_text(doc, "Band 3", font_size=11)
    add_blank_lines(doc, 2)

    teaser_paragraphs = [
        "Man kann eine Nummer nicht anrufen, wenn niemand mehr da ist, der abhebt.",
        "Das wusste Nora noch nicht. Sie wählte trotzdem.",
        "Vor ihr auf dem Küchentisch lag die Visitenkarte. Drei Tage zuvor hatte Nora sie im Hausflur gefunden, die Tinte noch feucht — als hätte jemand sie eine Minute vorher geschrieben. Jetzt war sie trocken. Nora drehte sie zwischen den Fingern. Vorderseite: *M. Silber.* Eine Telefonnummer. Rückseite: zwei Worte.",
        "*Nicht alleine.*",
        "Sie hatte die Nummer schon sechsmal gewählt. Heute war das siebte Mal.",
        "Es klingelte.",
        "\"Sie geht nicht ran\", sagte Theo.",
        "Er saß ihr gegenüber, das Kinn auf den Armen, und sah dem Handy beim Klingeln zu, als wäre es ein Tier, das gleich beißen könnte.",
        "\"Vielleicht beim nächsten Mal.\"",
        "\"Das hast du beim vierten Mal auch gesagt.\"",
        "Es klingelte ein viertes Mal. Ein fünftes.",
        "Und dann — beim sechsten — ein Knacken. Ganz kurz. Als hätte am anderen Ende jemand den Hörer abgenommen. Nora drückte das Handy fester ans Ohr.",
        "\"Hallo?\", sagte sie. \"Frau Silber?\"",
        "Nichts. Kein Atmen, keine Stimme. Nur ein leises Rauschen, weit weg, wie Wind in einem leeren Raum. Dann das Tuten wieder, gleichmäßig und kalt.",
        "Beim siebten Klingeln kam ein langer Ton.",
        "Nora legte auf. Ihre Hand war nicht ganz ruhig.",
    ]

    for i, para in enumerate(teaser_paragraphs):
        add_formatted_paragraph(doc, apply_typography(para), first_line_indent=(i > 0))

    add_blank_lines(doc, 2)
    add_centered_text(doc, "Jetzt überall erhältlich.", font_size=11, italic=True)

    # --- 3. Serien-Übersicht ---
    add_page_break(doc)
    add_blank_lines(doc, 2)
    add_centered_text(doc, "Die Geisterspürer — Alle Bände", font_size=14, bold=True)
    add_blank_lines(doc, 1)

    for band_num, title in [
        (1, "Das Haus, das flüstert"),
        (2, "Der Friedhof ohne Namen"),
        (3, "Schatten sieht mehr"),
        (4, "Die zugemauerte Tür"),
        (5, "Der Schleier"),
    ]:
        add_formatted_paragraph(
            doc,
            f"**Band {band_num}:** {title}",
            first_line_indent=False,
        )

    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Übrigens: Die Geisterspürer gibt es auch als interaktives Mitmach-Buch, "
        "in dem *du* entscheidest, was Nora und Theo als Nächstes tun — mit vielen "
        "Wegen und Enden. Frag mal danach!",
        first_line_indent=False,
    )

    # --- 4. Cross-Verweis andere Serien ---
    add_blank_lines(doc, 1)
    add_centered_text(doc, "* * *", font_size=11)
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Mehr spannende Abenteuer von Benjamin Krug:", font_size=14, bold=True)
    add_blank_lines(doc, 1)

    add_centered_text(doc, "Die Herrenhaus-Detektive", font_size=13, bold=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Niemand darf das alte Herrenhaus betreten. Aber Jonas, Mila und Ben "
        "finden einen Schlüssel — und hinter der verschlossenen "
        "Tür wartet ein Geheimnis, das seit dreißig Jahren niemand "
        "lüften durfte.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Spannendes Detektivabenteuer ab 8 Jahren.", font_size=11, italic=True)

    add_blank_lines(doc, 2)

    add_centered_text(doc, "Dein Fall – Du entscheidest!", font_size=13, bold=True)
    add_centered_text(doc, "Das verbotene Herrenhaus", font_size=11, italic=True)
    add_blank_lines(doc, 1)
    add_formatted_paragraph(
        doc,
        "Kein normales Buch: Hier bestimmst *du*, was passiert. An jeder wichtigen "
        "Stelle triffst du eine Entscheidung – und schlägst die Seite auf, die zu "
        "deiner Wahl gehört.",
        first_line_indent=False,
    )
    add_formatted_paragraph(
        doc,
        "Vier Wege durch das alte Herrenhaus, 14 verschiedene Enden. Findest du "
        "den richtigen? Oder tappst du in die Falle? Bei jedem Lesen wird die "
        "Geschichte neu.",
        first_line_indent=False,
    )
    add_blank_lines(doc, 1)
    add_centered_text(doc, "Interaktives Detektiv-Spielbuch ab 8 Jahren.", font_size=11, italic=True)


def add_chapter_heading(doc, title):
    """Add a chapter heading with page break, plus a bookmark for the TOC."""
    add_page_break(doc)
    add_blank_lines(doc, 3)
    p = doc.add_paragraph(title, style="Heading 1")
    for run in p.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(16)
        run.bold = True
    nummer = _kapitel_nummer(title)
    _add_bookmark(p, nummer, f"kapitel_{nummer}")
    add_blank_lines(doc, 2)


def add_scene_break(doc):
    add_blank_lines(doc, 1)
    add_centered_text(doc, "* * *", font_size=11)
    add_blank_lines(doc, 1)


def parse_and_build(doc, content):
    """Parse markdown content and build the DOCX."""
    match = re.search(r'^# Kapitel 1', content, re.MULTILINE)
    if not match:
        raise ValueError("Could not find '# Kapitel 1' in manuscript")
    body = content[match.start():]

    # Band 2 hat denselben einfachen ENDE-Marker wie Band 1 (kein Sonderfall
    # wie Band 5) -- verifiziert 2026-09-16 gegen Manuskript_Band2_Komplett.md.
    body = re.sub(r'\n---\n\n\*\*ENDE BAND 2\*\*\n\n---\n*', '', body)
    body = re.sub(r'\n\*\*ENDE BAND 2\*\*\n*', '', body)

    lines = body.split('\n')

    def _naechste_zeile_ist_kapitel(start):
        j = start
        while j < len(lines) and not lines[j].strip():
            j += 1
        return j < len(lines) and lines[j].strip().startswith('# Kapitel ')

    is_first_para_after_heading = False
    is_first_para_after_break = False
    i = 0

    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        if line.startswith('# Kapitel '):
            title = line[2:]
            add_chapter_heading(doc, title)
            is_first_para_after_heading = True
            i += 1
            continue

        if line == '---':
            if _naechste_zeile_ist_kapitel(i + 1):
                i += 1
                continue
            add_scene_break(doc)
            is_first_para_after_break = True
            i += 1
            continue

        no_indent = is_first_para_after_heading or is_first_para_after_break
        # Band 2 ist reines ASCII im Manuskript -> hier auf deutsche
        # Typografie umwandeln (Band 1 brauchte das nicht, weil das
        # Manuskript selbst schon korrigiert ist).
        text = apply_typography(line)
        add_formatted_paragraph(doc, text, first_line_indent=not no_indent)
        is_first_para_after_heading = False
        is_first_para_after_break = False
        i += 1


def main():
    with open(INPUT_FILE, "r", encoding="utf-8") as f:
        content = f.read()

    chapter_titles = re.findall(r"^# (Kapitel \d+.*)$", content, re.MULTILINE)
    if not chapter_titles:
        raise ValueError("Keine Kapitelueberschriften im Manuskript gefunden")

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    create_styles(doc)

    add_front_matter(doc, chapter_titles)
    parse_and_build(doc, content)
    add_back_matter(doc)

    doc.save(OUTPUT_FILE)
    print(f"{OUTPUT_FILE} erstellt!")

    para_count = len(doc.paragraphs)
    print(f"Absätze: {para_count}")


if __name__ == "__main__":
    main()
